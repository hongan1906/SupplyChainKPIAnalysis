"""
Generate a synthetic supply-chain dataset and save it as supply_chain_data.xlsx.

Change SEED (or any parameter below) to get a different dataset, then re-run
run_analysis.py to refresh the database, results and charts.
"""
import numpy as np
import pandas as pd
from datetime import date, timedelta

SEED = 42
N_PRODUCTS = 40
N_COMPONENTS = 60
N_CUSTOMERS = 25
N_VENDORS = 8
N_ORDERS = 1200
ZIPF_S = 1.3          # demand concentration: higher = fewer products dominate
START, END = date(2025, 1, 1), date(2025, 12, 31)

rng = np.random.default_rng(SEED)
days = (END - START).days


def rand_date(n):
    return [START + timedelta(days=int(d)) for d in rng.integers(0, days + 1, n)]


# ---------------- Master data ----------------
families = ["Controller", "Sensor", "Power Unit", "Valve", "Pump", "Display", "Gateway", "Actuator"]
products = pd.DataFrame({
    "product_id": [f"P{i:03d}" for i in range(1, N_PRODUCTS + 1)],
    "product_name": [f"{families[i % len(families)]} {chr(65 + i // len(families))}{i % 7 + 1}"
                     for i in range(N_PRODUCTS)],
    "product_family": [families[i % len(families)] for i in range(N_PRODUCTS)],
    "list_price": np.round(rng.uniform(40, 900, N_PRODUCTS), 2),
})
# Popularity follows a Zipf-like long tail (a few best sellers, many slow sellers)
popularity = rng.permutation(1.0 / np.arange(1, N_PRODUCTS + 1) ** ZIPF_S)
popularity = popularity / popularity.sum()

categories = ["Electronics", "Mechanical", "Packaging", "Fasteners", "Cables", "Plastics"]
components = pd.DataFrame({
    "component_id": [f"C{i:03d}" for i in range(1, N_COMPONENTS + 1)],
    "component_name": [f"{categories[i % len(categories)][:4].upper()}-{1000 + i * 7}"
                       for i in range(N_COMPONENTS)],
    "category": [categories[i % len(categories)] for i in range(N_COMPONENTS)],
    "unit_cost": np.round(rng.uniform(0.2, 12, N_COMPONENTS), 2),
})

countries = ["Finland", "Sweden", "Germany", "Estonia", "Poland", "Netherlands"]
customers = pd.DataFrame({
    "customer_id": [f"CU{i:03d}" for i in range(1, N_CUSTOMERS + 1)],
    "customer_name": [f"Customer {i:03d} Oy" for i in range(1, N_CUSTOMERS + 1)],
    "country": rng.choice(countries, N_CUSTOMERS),
})

vendors = pd.DataFrame({
    "vendor_id": [f"V{i:02d}" for i in range(1, N_VENDORS + 1)],
    "vendor_name": [f"Supplier {i:02d}" for i in range(1, N_VENDORS + 1)],
    "country": rng.choice(["Finland", "Germany", "China", "Poland", "Taiwan"], N_VENDORS),
    "standard_lead_time_days": rng.integers(7, 45, N_VENDORS),
})

# ---------------- Bill of materials ----------------
# Leave the last component unused by any product (shows up as "No activity")
usable = components["component_id"].iloc[:-1].tolist()
bom_rows = []
for pid in products["product_id"]:
    for cid in rng.choice(usable, rng.integers(3, 7), replace=False):
        bom_rows.append({"product_id": pid, "component_id": cid,
                         "quantity": int(rng.choice([1, 1, 1, 2, 2, 4, 6]))})
bom = pd.DataFrame(bom_rows)

# ---------------- Sales orders ----------------
so_products = rng.choice(products["product_id"], N_ORDERS, p=popularity)
price_map = products.set_index("product_id")["list_price"]
order_dates = rand_date(N_ORDERS)
sales_orders = pd.DataFrame({
    "order_id": [f"SO{i:05d}" for i in range(1, N_ORDERS + 1)],
    "order_date": order_dates,
    "customer_id": rng.choice(customers["customer_id"], N_ORDERS),
    "product_id": so_products,
    "quantity": rng.integers(1, 25, N_ORDERS),
    # Price varies around list price (discounts / surcharges)
    "unit_price": [round(price_map[p] * rng.uniform(0.85, 1.05), 2) for p in so_products],
    "requested_ship_date": [d + timedelta(days=int(rng.integers(5, 15))) for d in order_dates],
})
sales_orders["actual_ship_date"] = [
    r + timedelta(days=int(rng.choice([-2, -1, 0, 0, 0, 1, 2, 5, 9])))
    for r in sales_orders["requested_ship_date"]
]
sales_orders = sales_orders.sort_values("order_date").reset_index(drop=True)

# ---------------- Purchase orders ----------------
# Component demand implied by sales, so receipts can be over/under-bought on purpose
demand = (sales_orders.merge(bom, on="product_id", suffixes=("_so", "_bom"))
          .assign(used=lambda d: d["quantity_so"] * d["quantity_bom"])
          .groupby("component_id")["used"].sum())

# Buying behaviour per component: received = demand x factor
profile = rng.choice(["slow", "normal", "fast", "short"], N_COMPONENTS, p=[0.25, 0.40, 0.27, 0.08])
factor_range = {"slow": (5.2, 9), "normal": (1.8, 4.5), "fast": (1.05, 1.6), "short": (0.6, 0.95)}

po_headers, po_lines = [], []
po_no, line_no = 1, 1
no_po_component = components["component_id"].iloc[-2]   # has demand but never purchased
for i, cid in enumerate(components["component_id"]):
    if cid == no_po_component:
        continue
    d = demand.get(cid, 0)
    target = int(d * rng.uniform(*factor_range[profile[i]])) if d > 0 else int(rng.integers(50, 400))
    n_po = int(rng.integers(2, 7))
    splits = rng.dirichlet(np.ones(n_po)) * target
    vendor = rng.choice(vendors["vendor_id"])
    lead = int(vendors.set_index("vendor_id").loc[vendor, "standard_lead_time_days"])
    for qty in splits:
        qty = max(int(round(qty)), 1)
        po_date = rand_date(1)[0]
        promised = po_date + timedelta(days=lead)
        received = promised + timedelta(days=int(rng.choice([-3, 0, 0, 2, 4, 10, 18])))
        po_id = f"PO{po_no:05d}"
        po_headers.append({"po_id": po_id, "vendor_id": vendor, "po_date": po_date})
        ordered = int(qty * rng.choice([1.0, 1.0, 1.0, 1.05, 1.15]))
        po_lines.append({"po_line_id": f"PL{line_no:05d}", "po_id": po_id, "component_id": cid,
                         "ordered_qty": ordered, "delivered_qty": qty,
                         "unit_cost": float(components.loc[i, "unit_cost"]),
                         "promised_date": promised, "received_date": received})
        po_no += 1
        line_no += 1

purchase_orders = pd.DataFrame(po_headers).sort_values("po_date").reset_index(drop=True)
po_lines = pd.DataFrame(po_lines)

# ---------------- Data dictionary ----------------
dictionary = pd.DataFrame([
    ("products", "Finished goods master", "product_id, product_name, product_family, list_price"),
    ("components", "Component / raw material master", "component_id, component_name, category, unit_cost"),
    ("bom", "Bill of materials: component qty per 1 unit of product", "product_id, component_id, quantity"),
    ("customers", "Customer master", "customer_id, customer_name, country"),
    ("vendors", "Supplier master", "vendor_id, vendor_name, country, standard_lead_time_days"),
    ("sales_orders", "Demand: one row per order line (2025)",
     "order_id, order_date, customer_id, product_id, quantity, unit_price, requested_ship_date, actual_ship_date"),
    ("purchase_orders", "Supply: PO headers", "po_id, vendor_id, po_date"),
    ("po_lines", "Supply: PO lines with receipts",
     "po_line_id, po_id, component_id, ordered_qty, delivered_qty, unit_cost, promised_date, received_date"),
], columns=["table", "description", "columns"])

sheets = {
    "README": dictionary, "products": products, "components": components, "bom": bom,
    "customers": customers, "vendors": vendors, "sales_orders": sales_orders,
    "purchase_orders": purchase_orders, "po_lines": po_lines,
}

out = "supply_chain_data.xlsx"
with pd.ExcelWriter(out, engine="openpyxl", date_format="YYYY-MM-DD") as xw:
    for name, df in sheets.items():
        df.to_excel(xw, sheet_name=name, index=False)

# Formatting: Arial, bold header, frozen header row, filters, column widths
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = load_workbook(out)
header_fill = PatternFill("solid", start_color="1F4E78")
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=10)
    for c in ws[1]:
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = header_fill
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 70)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if hasattr(c.value, "year"):
                c.number_format = "yyyy-mm-dd"
ws = wb["README"]
ws.insert_rows(1, 3)
ws["A1"] = "Synthetic supply chain dataset (generated with generate_data.py, seed = %d)" % SEED
ws["A1"].font = Font(name="Arial", size=12, bold=True)
ws["A2"] = "All data is fictional. Edit any sheet or re-run the generator, then run run_analysis.py to refresh results."
ws["A2"].font = Font(name="Arial", size=10, italic=True)
ws.freeze_panes = None
ws.auto_filter.ref = None
wb.save(out)

for name, df in sheets.items():
    print(f"{name:16s} {len(df):5d} rows")
